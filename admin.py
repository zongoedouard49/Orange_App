import io
from datetime import timedelta, datetime
from typing import Optional

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from fastapi import APIRouter, Depends, HTTPException, Query, Request, status
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

import auth
import models
import schemas
from config import ACCESS_TOKEN_EXPIRE_MINUTES
from database import get_db

router = APIRouter(prefix="/admin", tags=["Admin"])


@router.post(
    "/register",
    response_model=dict,
    status_code=status.HTTP_201_CREATED,
    summary="Créer un compte administrateur",
)
def register_admin(
    admin_data: schemas.AdminCreate,
    db: Session = Depends(get_db),
):
    existing = (
        db.query(models.Admin)
        .filter(models.Admin.cuid == admin_data.cuid)
        .first()
    )
    if existing:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=f"Un administrateur avec le CUID '{admin_data.cuid}' existe déjà.",
        )

    new_admin = models.Admin(
        cuid=admin_data.cuid,
        password=auth.hash_password(admin_data.password),
    )
    db.add(new_admin)
    db.commit()
    db.refresh(new_admin)

    return {
        "message": f"Administrateur '{new_admin.cuid}' créé avec succès. Le compte est inactif par défaut — un administrateur doit l'activer.",
        "id": new_admin.id,
        "actif": new_admin.actif,
    }


@router.post(
    "/login",
    response_model=schemas.Token,
    summary="Authentification administrateur",
)
def login_admin(
    credentials: schemas.AdminLogin,
    request: Request,
    db: Session = Depends(get_db),
):
    admin = (
        db.query(models.Admin)
        .filter(models.Admin.cuid == credentials.cuid)
        .first()
    )

    if not admin or not auth.verify_password(credentials.password, admin.password):
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Nom d'utilisateur ou mot de passe incorrect.",
        )

    if not admin.actif:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Votre compte est inactif. Veuillez contacter un administrateur pour l'activer.",
        )

    token = auth.create_access_token(
        data={"sub": admin.cuid},
        expires_delta=timedelta(minutes=ACCESS_TOKEN_EXPIRE_MINUTES),
    )

    # Log admin login
    client_ip = request.client.host if request.client else "unknown"
    db.add(models.LoginLog(cuid=admin.cuid, action="login_admin", adresse_ip=client_ip))
    db.commit()

    return {"access_token": token, "token_type": "bearer"}


@router.get("/comptes", summary="Lister les comptes administrateurs")
def list_admins(
    current_admin: models.Admin = Depends(auth.get_current_admin),
    db: Session = Depends(get_db),
):
    return [
        {
            "id": a.id,
            "cuid": a.cuid,
            "actif": a.actif,
            "created_at": a.created_at.isoformat() if a.created_at else None,
        }
        for a in db.query(models.Admin).order_by(models.Admin.id).all()
    ]


@router.patch("/comptes/{admin_id}/activation", summary="Activer/désactiver un compte")
def set_admin_status(
    admin_id: int,
    actif: bool,
    current_admin: models.Admin = Depends(auth.get_current_admin),
    db: Session = Depends(get_db),
):
    target = db.query(models.Admin).filter(models.Admin.id == admin_id).first()
    if target is None:
        raise HTTPException(status_code=404, detail="Administrateur introuvable.")
    if target.id == current_admin.id and not actif:
        raise HTTPException(status_code=400, detail="Vous ne pouvez pas désactiver votre propre compte.")
    target.actif = actif
    db.commit()
    return {"message": f"Compte '{target.cuid}' {'activé' if actif else 'désactivé'}.", "id": target.id, "actif": target.actif}


@router.get("/blocked-ips", summary="Lister les IP bloquées")
def list_blocked_ips(
    current_admin: models.Admin = Depends(auth.get_current_admin),
    db: Session = Depends(get_db),
):
    records = db.query(models.BlockedIP).order_by(models.BlockedIP.id.desc()).all()
    return [
        {
            "id": r.id,
            "adresse_ip": r.adresse_ip,
            "cuid": r.cuid or "",
            "adresse_mac": r.adresse_mac or "",
            "raison": r.raison,
            "date_creation": r.date_creation.isoformat() if r.date_creation else "",
        }
        for r in records
    ]


@router.delete("/blocked-ips/{ip_id}", summary="Débloquer une IP")
def unblock_ip(
    ip_id: int,
    current_admin: models.Admin = Depends(auth.get_current_admin),
    db: Session = Depends(get_db),
):
    record = db.query(models.BlockedIP).filter(models.BlockedIP.id == ip_id).first()
    if not record:
        raise HTTPException(status_code=404, detail="IP bloquée introuvable.")
    ip_addr = record.adresse_ip
    db.delete(record)
    db.commit()
    # Also remove from in-memory set
    from main import _blocked_ips
    _blocked_ips.discard(ip_addr)
    return {"message": f"L'adresse IP '{ip_addr}' a été débloquée."}


@router.post("/logout", summary="Déconnexion administrateur (log)")
def logout_admin(
    request: Request,
    current_admin: models.Admin = Depends(auth.get_current_admin),
    db: Session = Depends(get_db),
):
    client_ip = request.client.host if request.client else "unknown"
    db.add(models.LoginLog(cuid=current_admin.cuid, action="logout_admin", adresse_ip=client_ip))
    db.commit()
    return {"message": "Déconnexion enregistrée."}


# ─── Logs de connexion ──────────────────────────────────────────

@router.get("/logs", summary="Lister les logs de connexion")
def list_logs(
    cuid: Optional[str] = Query(None),
    adresse_mac: Optional[str] = Query(None),
    page: int = Query(1, ge=1),
    per_page: int = Query(20, ge=1, le=100),
    current_admin: models.Admin = Depends(auth.get_current_admin),
    db: Session = Depends(get_db),
):
    query = db.query(models.LoginLog)
    if cuid:
        query = query.filter(models.LoginLog.cuid.ilike(f"%{cuid}%"))
    if adresse_mac:
        query = query.filter(models.LoginLog.adresse_mac.ilike(f"%{adresse_mac}%"))

    total = query.count()
    records = query.order_by(models.LoginLog.id.desc()).offset((page - 1) * per_page).limit(per_page).all()

    return {
        "data": [
            {
                "id": r.id,
                "cuid": r.cuid,
                "action": getattr(r, "action", "login") or "login",
                "adresse_ip": r.adresse_ip or "",
                "adresse_mac": r.adresse_mac or "",
                "date_creation": r.date_creation.isoformat() if r.date_creation else "",
            }
            for r in records
        ],
        "total": total,
        "page": page,
        "per_page": per_page,
        "total_pages": max(1, (total + per_page - 1) // per_page),
    }


@router.get("/logs/exporter", summary="Exporter les logs en Excel")
def exporter_logs(
    cuid: Optional[str] = Query(None),
    adresse_mac: Optional[str] = Query(None),
    current_admin: models.Admin = Depends(auth.get_current_admin),
    db: Session = Depends(get_db),
):
    query = db.query(models.LoginLog)
    if cuid:
        query = query.filter(models.LoginLog.cuid.ilike(f"%{cuid}%"))
    if adresse_mac:
        query = query.filter(models.LoginLog.adresse_mac.ilike(f"%{adresse_mac}%"))

    records = query.order_by(models.LoginLog.id.desc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Logs de connexion"

    BLEU_FONCE = "1F4E79"
    BLANC = "FFFFFF"
    en_tete_font = Font(bold=True, color=BLANC, size=11, name="Calibri")
    en_tete_fill = PatternFill(start_color=BLEU_FONCE, end_color=BLEU_FONCE, fill_type="solid")
    centre = Alignment(horizontal="center", vertical="center")
    bordure = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    colonnes = [("ID", 8), ("CUID", 20), ("Action", 18), ("Adresse IP", 20), ("Adresse MAC", 22), ("Date de connexion", 25)]
    for col_idx, (titre, largeur) in enumerate(colonnes, start=1):
        cell = ws.cell(row=1, column=col_idx, value=titre)
        cell.font = en_tete_font
        cell.fill = en_tete_fill
        cell.alignment = centre
        cell.border = bordure
        ws.column_dimensions[cell.column_letter].width = largeur

    ws.freeze_panes = "A2"

    for row_idx, rec in enumerate(records, start=2):
        valeurs = [rec.id, rec.cuid, getattr(rec, "action", "login") or "login",
                   rec.adresse_ip or "", rec.adresse_mac or "",
                   rec.date_creation.strftime("%d/%m/%Y %H:%M:%S") if rec.date_creation else ""]
        for col_idx, valeur in enumerate(valeurs, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=valeur)
            cell.border = bordure
            cell.alignment = centre

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"logs_connexion_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
