import json
import os

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, status

import auth
import models
import realtime
import schemas
from config import SERVICES_FILE, STATUTS_VALIDES

router = APIRouter(prefix="/services", tags=["Services"])


def _load() -> dict:
    with open(SERVICES_FILE, "r", encoding="utf-8") as f:
        data = json.load(f)
    data.setdefault("services", [])
    data.setdefault("plateformes", [])
    data.setdefault("partenaires", [])
    return data


def _save(data: dict) -> None:
    os.makedirs(os.path.dirname(SERVICES_FILE), exist_ok=True)
    with open(SERVICES_FILE, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    realtime.publish("services")  # temps réel : listes mises à jour chez les clients connectés


@router.get("/", summary="Lister tout")
def get_all(current_admin: models.Admin = Depends(auth.get_current_user),):
    data = _load()
    data["statuts"] = sorted(STATUTS_VALIDES)
    return data


@router.get("/plateformes", summary="Lister les plateformes")
def get_plateformes(current_admin: models.Admin = Depends(auth.get_current_user),):
    return {"plateformes": _load().get("plateformes", [])}


@router.get("/partenaires", summary="Lister les partenaires")
def get_partenaires(current_admin: models.Admin = Depends(auth.get_current_user),):
    return {"partenaires": _load().get("partenaires", [])}


@router.post("/update", summary="Remplacer la liste (admin)")
def update_liste(
    body: schemas.ServicesUpdate,
    current_admin: models.Admin = Depends(auth.get_current_admin),
):
    seen = set()
    unique = []
    for s in body.name:
        key = s.lower()
        if key not in seen:
            seen.add(key)
            unique.append(s)
    data = _load()
    data["services"] = unique
    _save(data)
    return data


@router.post("/delete", summary="Supprimer un élément")
def supprimer(body: schemas.ServiceDeleteRequest,current_admin: models.Admin = Depends(auth.get_current_admin),):
    data = _load()
    type_ = body.type_
    items = data.get(type_, [])

    match = next(
        (s for s in items if s.lower() == body.name.lower()),
        None,
    )
    if not match:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"'{body.name}' introuvable dans '{type_}'.",
        )

    items.remove(match)
    data[type_] = items
    _save(data)
    return data


@router.get("/modele-excel", summary="Télécharger l'exemplaire Excel pour l'import des listes")
def modele_excel_services(current_admin: models.Admin = Depends(auth.get_current_admin)):
    import io
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Listes"

    ORANGE = "FF7900"
    BLANC = "FFFFFF"
    en_tete_font = Font(bold=True, color=BLANC, size=11, name="Calibri")
    en_tete_fill = PatternFill(start_color=ORANGE, end_color=ORANGE, fill_type="solid")
    centre = Alignment(horizontal="center", vertical="center")
    bordure = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    entetes = ["Plateformes", "Services", "Partenaires"]
    for col_idx, titre in enumerate(entetes, start=1):
        cell = ws.cell(row=1, column=col_idx, value=titre)
        cell.font = en_tete_font
        cell.fill = en_tete_fill
        cell.alignment = centre
        cell.border = bordure
        ws.column_dimensions[cell.column_letter].width = 28

    exemples = [
        ["WhatsApp", "Direction Informatique", "Entreprise ACME"],
        ["Teams", "Direction Ressources Humaines", "Partenaire XYZ"],
    ]
    for row_idx, ligne in enumerate(exemples, start=2):
        for col_idx, valeur in enumerate(ligne, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=valeur)
            cell.border = bordure
            cell.alignment = Alignment(horizontal="left", vertical="center")

    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    from fastapi.responses import StreamingResponse
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=modele_import_listes.xlsx"},
    )


@router.post("/importer-excel", summary="Importer depuis un fichier Excel")
async def importer_excel(
    file: UploadFile = File(...),
    current_admin: models.Admin = Depends(auth.get_current_admin),
):
    if not file.filename or not file.filename.lower().endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="Veuillez fournir un fichier Excel (.xlsx).")

    try:
        import io
        import openpyxl
        contenu = await file.read()
        wb = openpyxl.load_workbook(io.BytesIO(contenu), data_only=True)
        ws = wb.active
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Fichier Excel illisible ou invalide") from exc

    def normaliser(valeur):
        return " ".join(str(valeur).split()).strip() if valeur is not None else ""

    premiere_ligne = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not premiere_ligne:
        raise HTTPException(status_code=400, detail="Le fichier Excel est vide.")

    entetes = [normaliser(c).casefold() for c in premiere_ligne]
    colonnes_attendues = {"plateformes": "Plateformes", "services": "Services", "partenaires": "Partenaires"}
    colonnes = {}
    for cle in colonnes_attendues:
        if cle in entetes:
            colonnes[cle] = entetes.index(cle)

    manquantes = [libelle for cle, libelle in colonnes_attendues.items() if cle not in colonnes]
    if len(manquantes) == len(colonnes_attendues):
        raise HTTPException(
            status_code=400,
            detail=f"Le fichier ne contient aucune des colonnes attendues : {', '.join(colonnes_attendues.values())}.",
        )

    valeurs_plateformes, valeurs_services, valeurs_partenaires = [], [], []
    dest_map = {
        "plateformes": valeurs_plateformes,
        "services": valeurs_services,
        "partenaires": valeurs_partenaires,
    }
    for ligne in ws.iter_rows(min_row=2, values_only=True):
        if ligne is None:
            continue
        for cle, destination in dest_map.items():
            if cle in colonnes:
                idx = colonnes[cle]
                if idx < len(ligne):
                    valeur = normaliser(ligne[idx])
                    if valeur:
                        destination.append(valeur)

    data = _load()

    def ajouter_uniques(cle, valeurs):
        existants = data.setdefault(cle, [])
        vus = {v.strip().casefold() for v in existants if isinstance(v, str)}
        ajoutes = []
        for valeur in valeurs:
            if valeur.casefold() not in vus:
                existants.append(valeur)
                vus.add(valeur.casefold())
                ajoutes.append(valeur)
        return ajoutes

    np = ajouter_uniques("plateformes", valeurs_plateformes)
    ns = ajouter_uniques("services", valeurs_services)
    npart = ajouter_uniques("partenaires", valeurs_partenaires)
    _save(data)
    message = "Import terminé."
    if manquantes:
        message += f" Colonne(s) absente(s) et ignorée(s) : {', '.join(manquantes)}."
    return {
        "message": message,
        "colonnes_manquantes": manquantes,
        "plateformes_ajoutees": np,
        "services_ajoutes": ns,
        "partenaires_ajoutes": npart,
    }


@router.post("/ajout")
def ajouter_valeur(data: schemas.AjouterValeurRequest,current_admin: models.Admin = Depends(auth.get_current_admin),):
    try:
        with open(SERVICES_FILE, "r", encoding="utf-8") as f:
            contenu = json.load(f)
    except json.JSONDecodeError:
        raise HTTPException(status_code=500, detail="Le fichier JSON est invalide.")

    cle = f'{data.cle.strip()}s'
    valeur = data.valeur.strip()

    if not cle:
        raise HTTPException(status_code=400, detail="La clé est obligatoire.")
    if not valeur:
        raise HTTPException(status_code=400, detail="La valeur est obligatoire.")

    if cle not in contenu:
        contenu[cle] = []

    if not isinstance(contenu[cle], list):
        raise HTTPException(status_code=400, detail=f"La clé '{cle}' ne contient pas une liste.")

    if valeur.lower() in [v.lower() for v in contenu[cle]]:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=f"La valeur '{valeur}' existe déjà dans '{cle}'."
        )

    contenu[cle].append(valeur)

    with open(SERVICES_FILE, "w", encoding="utf-8") as f:
        json.dump(contenu, f, ensure_ascii=False, indent=4)
    realtime.publish("services")

    return {
        "success": True,
        "message": "Valeur ajoutée avec succès.",
        "cle": cle,
        "valeur": valeur,
        "ajoute": True
    }
