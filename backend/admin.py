import io
from datetime import timedelta, datetime
from typing import List, Optional

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from fastapi import APIRouter, Depends, File, HTTPException, Query, Request, UploadFile, status
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy import func
from sqlalchemy.orm import Session

import auth
import models
import schemas
from config import ACCESS_TOKEN_EXPIRE_MINUTES, DEFAULT_PAGE_SIZE, MAX_PAGE_SIZE
from database import get_db

router = APIRouter(prefix="/admin", tags=["Admin"])


@router.post(
    "/register",
    response_model=dict,
    status_code=status.HTTP_201_CREATED,
    summary="Créer un compte administrateur",
)
def register_admin(
    admin_data: schemas.AdminCreate,
    db: Session = Depends(get_db),
):
    existing = (
        db.query(models.Admin)
        .filter(models.Admin.cuid == admin_data.cuid)
        .first()
    )
    if existing:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=f"Un administrateur avec le CUID '{admin_data.cuid}' existe déjà.",
        )

    new_admin = models.Admin(
        cuid=admin_data.cuid,
        password=auth.hash_password(admin_data.password),
    )
    db.add(new_admin)
    db.commit()
    db.refresh(new_admin)

    return {
        "message": f"Administrateur '{new_admin.cuid}' créé avec succès. Le compte est inactif par défaut — un administrateur doit l'activer.",
        "id": new_admin.id,
        "actif": new_admin.actif,
    }


@router.post(
    "/login",
    response_model=schemas.Token,
    summary="Authentification administrateur",
)
def login_admin(
    credentials: schemas.AdminLogin,
    request: Request,
    db: Session = Depends(get_db),
):
    admin = (
        db.query(models.Admin)
        .filter(models.Admin.cuid == credentials.cuid)
        .first()
    )

    if not admin or not auth.verify_password(credentials.password, admin.password):
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Nom d'utilisateur ou mot de passe incorrect.",
        )

    if not admin.actif:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Votre compte est inactif. Veuillez contacter un administrateur pour l'activer.",
        )

    token = auth.create_access_token(
        data={"sub": admin.cuid},
        expires_delta=timedelta(minutes=ACCESS_TOKEN_EXPIRE_MINUTES),
    )

    # Log admin login
    client_ip = request.client.host if request.client else "unknown"
    db.add(models.LoginLog(cuid=admin.cuid, action="login_admin", adresse_ip=client_ip))
    db.commit()

    return {"access_token": token, "token_type": "bearer"}


@router.get("/comptes", summary="Lister les comptes administrateurs")
def list_admins(
    current_admin: models.Admin = Depends(auth.get_current_admin),
    db: Session = Depends(get_db),
):
    return [
        {
            "id": a.id,
            "cuid": a.cuid,
            "actif": a.actif,
            "created_at": a.created_at.isoformat() if a.created_at else None,
        }
        for a in db.query(models.Admin).order_by(models.Admin.id).all()
    ]


@router.patch("/comptes/{admin_id}/activation", summary="Activer/désactiver un compte")
def set_admin_status(
    admin_id: int,
    actif: bool,
    current_admin: models.Admin = Depends(auth.get_current_admin),
    db: Session = Depends(get_db),
):
    target = db.query(models.Admin).filter(models.Admin.id == admin_id).first()
    if target is None:
        raise HTTPException(status_code=404, detail="Administrateur introuvable.")
    if target.id == current_admin.id and not actif:
        raise HTTPException(status_code=400, detail="Vous ne pouvez pas désactiver votre propre compte.")
    target.actif = actif
    db.commit()
    return {"message": f"Compte '{target.cuid}' {'activé' if actif else 'désactivé'}.", "id": target.id, "actif": target.actif}


@router.get("/blocked-ips", summary="Lister les IP bloquées")
def list_blocked_ips(
    current_admin: models.Admin = Depends(auth.get_current_admin),
    db: Session = Depends(get_db),
):
    records = db.query(models.BlockedIP).order_by(models.BlockedIP.id.desc()).all()
    return [
        {
            "id": r.id,
            "adresse_ip": r.adresse_ip,
            "cuid": r.cuid or "",
            "adresse_mac": r.adresse_mac or "",
            "raison": r.raison,
            "date_creation": r.date_creation.isoformat() if r.date_creation else "",
        }
        for r in records
    ]


@router.delete("/blocked-ips/{ip_id}", summary="Débloquer une IP")
def unblock_ip(
    ip_id: int,
    current_admin: models.Admin = Depends(auth.get_current_admin),
    db: Session = Depends(get_db),
):
    record = db.query(models.BlockedIP).filter(models.BlockedIP.id == ip_id).first()
    if not record:
        raise HTTPException(status_code=404, detail="IP bloquée introuvable.")
    ip_addr = record.adresse_ip
    db.delete(record)
    db.commit()
    # Also remove from in-memory set
    from main import _blocked_ips
    _blocked_ips.discard(ip_addr)
    return {"message": f"L'adresse IP '{ip_addr}' a été débloquée."}


@router.post("/logout", summary="Déconnexion administrateur (log)")
def logout_admin(
    request: Request,
    current_admin: models.Admin = Depends(auth.get_current_admin),
    db: Session = Depends(get_db),
):
    client_ip = request.client.host if request.client else "unknown"
    db.add(models.LoginLog(cuid=current_admin.cuid, action="logout_admin", adresse_ip=client_ip))
    db.commit()
    return {"message": "Déconnexion enregistrée."}


# ─── Logs de connexion ──────────────────────────────────────────

@router.get("/logs", summary="Lister les logs de connexion")
def list_logs(
    cuid: Optional[str] = Query(None),
    adresse_mac: Optional[str] = Query(None),
    page: int = Query(1, ge=1),
    per_page: int = Query(DEFAULT_PAGE_SIZE, ge=1, le=MAX_PAGE_SIZE),
    current_admin: models.Admin = Depends(auth.get_current_admin),
    db: Session = Depends(get_db),
):
    query = db.query(models.LoginLog)
    if cuid:
        query = query.filter(models.LoginLog.cuid.ilike(f"%{cuid}%"))
    if adresse_mac:
        query = query.filter(models.LoginLog.adresse_mac.ilike(f"%{adresse_mac}%"))

    total = query.count()
    records = query.order_by(models.LoginLog.id.desc()).offset((page - 1) * per_page).limit(per_page).all()

    return {
        "data": [
            {
                "id": r.id,
                "cuid": r.cuid,
                "action": getattr(r, "action", "login") or "login",
                "adresse_ip": r.adresse_ip or "",
                "adresse_mac": r.adresse_mac or "",
                "date_creation": r.date_creation.isoformat() if r.date_creation else "",
            }
            for r in records
        ],
        "total": total,
        "page": page,
        "per_page": per_page,
        "total_pages": max(1, (total + per_page - 1) // per_page),
    }


@router.get("/logs/exporter", summary="Exporter les logs en Excel")
def exporter_logs(
    cuid: Optional[str] = Query(None),
    adresse_mac: Optional[str] = Query(None),
    current_admin: models.Admin = Depends(auth.get_current_admin),
    db: Session = Depends(get_db),
):
    query = db.query(models.LoginLog)
    if cuid:
        query = query.filter(models.LoginLog.cuid.ilike(f"%{cuid}%"))
    if adresse_mac:
        query = query.filter(models.LoginLog.adresse_mac.ilike(f"%{adresse_mac}%"))

    records = query.order_by(models.LoginLog.id.desc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Logs de connexion"

    BLEU_FONCE = "1F4E79"
    BLANC = "FFFFFF"
    en_tete_font = Font(bold=True, color=BLANC, size=11, name="Calibri")
    en_tete_fill = PatternFill(start_color=BLEU_FONCE, end_color=BLEU_FONCE, fill_type="solid")
    centre = Alignment(horizontal="center", vertical="center")
    bordure = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    colonnes = [("ID", 8), ("CUID", 20), ("Action", 18), ("Adresse IP", 20), ("Adresse MAC", 22), ("Date de connexion", 25)]
    for col_idx, (titre, largeur) in enumerate(colonnes, start=1):
        cell = ws.cell(row=1, column=col_idx, value=titre)
        cell.font = en_tete_font
        cell.fill = en_tete_fill
        cell.alignment = centre
        cell.border = bordure
        ws.column_dimensions[cell.column_letter].width = largeur

    ws.freeze_panes = "A2"

    for row_idx, rec in enumerate(records, start=2):
        valeurs = [rec.id, rec.cuid, getattr(rec, "action", "login") or "login",
                   rec.adresse_ip or "", rec.adresse_mac or "",
                   rec.date_creation.strftime("%d/%m/%Y %H:%M:%S") if rec.date_creation else ""]
        for col_idx, valeur in enumerate(valeurs, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=valeur)
            cell.border = bordure
            cell.alignment = centre

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"logs_connexion_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


# ══════════════════════════════════════════════════════════════════════════
# AJOUTER DES UTILISATEURS (import Excel en masse)
# ══════════════════════════════════════════════════════════════════════════

COLONNES_IMPORT = ["cuid", "nom", "prenom", "statut", "direction", "entreprise_partenaire", "plateforme", "description"]
LIBELLES_IMPORT = {
    "cuid": "CUID",
    "nom": "Nom",
    "prenom": "Prénom",
    "statut": "Statut",
    "direction": "Direction",
    "entreprise_partenaire": "Entreprise Partenaire",
    "plateforme": "Plateforme utilisée",
    "description": "Description de la plateforme",
}


def _champs_manquants_ligne(valeurs: dict) -> list:
    manquants = []
    for champ in COLONNES_IMPORT:
        valeur = (valeurs.get(champ) or "").strip()
        if champ == "entreprise_partenaire":
            # Obligatoire uniquement si le statut n'est pas "Employé"
            if valeurs.get("statut", "").strip().lower() != "employé" and not valeur:
                manquants.append(LIBELLES_IMPORT[champ])
            continue
        if not valeur:
            manquants.append(LIBELLES_IMPORT[champ])
    return manquants


def _cuids_deja_renseignes(db: Session, cuids) -> set:
    """Retourne (en MAJUSCULES) les CUID de la liste qui existent déjà dans la table renseignements."""
    a_verifier = list({c.strip().upper() for c in cuids if c and c.strip()})
    trouves = set()
    for i in range(0, len(a_verifier), 500):  # par lots (limite de variables SQLite)
        lot = a_verifier[i:i + 500]
        rows = db.query(func.upper(models.Renseignement.CUID)).filter(
            func.upper(models.Renseignement.CUID).in_(lot)
        ).all()
        trouves.update(r[0] for r in rows)
    return trouves


@router.get("/utilisateurs/modele-excel", summary="Télécharger le modèle Excel d'import des utilisateurs")
def modele_excel_utilisateurs(current_admin: models.Admin = Depends(auth.get_current_admin)):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Utilisateurs"

    ORANGE = "FF7900"
    BLANC = "FFFFFF"
    en_tete_font = Font(bold=True, color=BLANC, size=11, name="Calibri")
    en_tete_fill = PatternFill(start_color=ORANGE, end_color=ORANGE, fill_type="solid")
    centre = Alignment(horizontal="center", vertical="center", wrap_text=True)
    bordure = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    entetes = [LIBELLES_IMPORT[c] for c in COLONNES_IMPORT]
    largeurs = [16, 20, 20, 14, 22, 22, 24, 45]
    for col_idx, (titre, largeur) in enumerate(zip(entetes, largeurs), start=1):
        cell = ws.cell(row=1, column=col_idx, value=titre)
        cell.font = en_tete_font
        cell.fill = en_tete_fill
        cell.alignment = centre
        cell.border = bordure
        ws.column_dimensions[cell.column_letter].width = largeur
    ws.freeze_panes = "A2"

    exemple = [
        "J12345", "Dupont", "Jean", "Employé", "Direction Informatique",
        "", "WhatsApp", "Utilisation professionnelle pour la communication d'équipe",
    ]
    for col_idx, valeur in enumerate(exemple, start=1):
        cell = ws.cell(row=2, column=col_idx, value=valeur)
        cell.border = bordure
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=modele_import_utilisateurs.xlsx"},
    )


@router.post("/utilisateurs/verifier-import", summary="Vérifier un fichier Excel d'import avant envoi en base")
async def verifier_import_utilisateurs(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_admin: models.Admin = Depends(auth.get_current_admin),
):
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Veuillez fournir un fichier Excel (.xlsx).")

    contenu = await file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(contenu), data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="Le fichier fourni n'est pas un fichier Excel valide.")

    ws = wb.active
    lignes = []
    nb_erreurs = 0

    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue  # ligne complètement vide : ignorée

        row = list(row) + [None] * (len(COLONNES_IMPORT) - len(row))
        valeurs = {
            champ: (str(row[i]).strip() if row[i] is not None else "")
            for i, champ in enumerate(COLONNES_IMPORT)
        }
        champs_manquants = _champs_manquants_ligne(valeurs)
        if champs_manquants:
            nb_erreurs += 1
        lignes.append({"ligne": idx, **valeurs, "champs_manquants": champs_manquants})

    if not lignes:
        raise HTTPException(status_code=400, detail="Le fichier ne contient aucune donnée à importer.")

    # CUID déjà renseignés en base : signalés et ignorés (jamais ré-inscrits).
    # Plusieurs lignes du même CUID dans le fichier restent autorisées (une ligne par plateforme).
    existants = _cuids_deja_renseignes(db, (l["cuid"] for l in lignes))
    nb_doublons = 0
    for l in lignes:
        l["deja_inscrit"] = l["cuid"].strip().upper() in existants
        if l["deja_inscrit"]:
            nb_doublons += 1
            l["champs_manquants"] = []  # inutile de corriger une ligne qui sera ignorée
    nb_erreurs = sum(1 for l in lignes if l["champs_manquants"])

    return {"lignes": lignes, "total": len(lignes), "nb_erreurs": nb_erreurs, "nb_doublons": nb_doublons}


class LigneImportUtilisateur(BaseModel):
    cuid: str = ""
    nom: str = ""
    prenom: str = ""
    statut: str = ""
    direction: str = ""
    entreprise_partenaire: Optional[str] = ""
    plateforme: str = ""
    description: str = ""


class ImporterUtilisateursRequest(BaseModel):
    lignes: List[LigneImportUtilisateur]


@router.post("/utilisateurs/importer", status_code=status.HTTP_201_CREATED, summary="Importer en base les utilisateurs validés")
def importer_utilisateurs(
    data: ImporterUtilisateursRequest,
    db: Session = Depends(get_db),
    current_admin: models.Admin = Depends(auth.get_current_admin),
):
    if not data.lignes:
        raise HTTPException(status_code=400, detail="Aucune ligne à importer.")

    erreurs = []
    for i, l in enumerate(data.lignes, start=1):
        manquants = _champs_manquants_ligne(l.dict())
        if manquants:
            erreurs.append({"ligne": i, "champs_manquants": manquants})

    if erreurs:
        raise HTTPException(
            status_code=422,
            detail={"message": "Certaines lignes contiennent des champs vides. Veuillez corriger avant d'envoyer.", "erreurs": erreurs},
        )

    # Re-vérification côté serveur (le fichier a pu être analysé avant qu'un autre import n'ait eu lieu)
    existants = _cuids_deja_renseignes(db, (l.cuid for l in data.lignes))
    ignores = sorted({l.cuid.strip() for l in data.lignes if l.cuid.strip().upper() in existants})

    inseres = 0
    for l in data.lignes:
        if l.cuid.strip().upper() in existants:
            continue  # déjà inscrit : signalé, pas de nouvelle inscription
        entry = models.Renseignement(
            CUID=l.cuid.strip(),
            Nom=l.nom.strip(),
            prenom=l.prenom.strip(),
            statut=l.statut.strip(),
            entreprise_partenaire=(l.entreprise_partenaire or "").strip() or None,
            direction=l.direction.strip(),
            plateforme=l.plateforme.strip(),
            description=l.description.strip(),
        )
        db.add(entry)
        inseres += 1
    db.commit()

    message = f"{inseres} ligne(s) importée(s) avec succès dans la base de données."
    if ignores:
        message += f" {len(ignores)} CUID déjà renseigné(s) ignoré(s) : {', '.join(ignores)}."
    return {"message": message, "inseres": inseres, "ignores": ignores}
