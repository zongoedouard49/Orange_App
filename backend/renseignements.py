import io
import json
from datetime import datetime
from typing import Optional, List

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from fastapi import APIRouter, Depends, HTTPException, Query, Request, status
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

import auth
import models
import schemas
from config import DEFAULT_PAGE_SIZE, MAX_PAGE_SIZE
from database import get_db

router = APIRouter(prefix="/renseignements", tags=["Renseignements"],include_in_schema=False)


@router.post(
    "/enregistrer",
    status_code=status.HTTP_201_CREATED,
    summary="Enregistrer un nouveau renseignement (une ligne par plateforme)",
)
def enregistrer_renseignement(
    data: schemas.RenseignementCreate,
    request: Request,
    db: Session = Depends(get_db),
    current_user: dict = Depends(auth.get_current_user),
):
    # Vérifier si ce CUID a déjà soumis
    deja = db.query(models.Renseignement).filter(models.Renseignement.CUID == data.cuid).first()
    if deja:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Ce CUID a déjà soumis ses renseignements. Vous ne pouvez pas soumettre à nouveau.",
        )

    # Vérifier si ce CUID est bloqué
    bloque = db.query(models.BlockedAccount).filter(models.BlockedAccount.cuid == data.cuid).first()
    if bloque:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Votre compte a été bloqué. Veuillez contacter l'administrateur.",
        )

    created: list[dict] = []
    for pf in data.plateformes:
        entry = models.Renseignement(
            CUID=data.cuid,
            Nom=data.nom,
            prenom=data.prenom,
            statut=data.statut,
            entreprise_partenaire=data.entreprise_partenaire,
            direction=data.direction,
            poste=data.poste,
            referent=data.referent,
            plateforme=pf.text.strip(),
            description=pf.description.strip(),
        )
        db.add(entry)
        db.flush()
        created.append({
            "id": entry.id,
            "CUID": entry.CUID,
            "Nom": entry.Nom,
            "prenom": entry.prenom,
            "statut": entry.statut,
            "entreprise_partenaire": entry.entreprise_partenaire or "",
            "direction": entry.direction,
            "poste": entry.poste or "",
            "referent": entry.referent or "",
            "plateforme": entry.plateforme or "",
            "description": entry.description,
        })

    # Log the submission — récupération MAC via ARP
    client_ip = request.client.host if request.client else "unknown"
    from users import get_mac_from_arp
    mac_address = get_mac_from_arp(client_ip)
    db.add(models.RenseignementLog(
        cuid=data.cuid,
        adresse_ip=client_ip,
        adresse_mac=mac_address,
    ))
    db.commit()
    return {"message": f"{len(created)} ligne(s) enregistrée(s) avec succès.", "data": created}


@router.get("/list", summary="Lister les renseignements avec filtres et pagination")
def lister_renseignements(
    direction: Optional[str] = Query(None),
    statut: Optional[str] = Query(None),
    cuid: Optional[str] = Query(None),
    plateforme: Optional[str] = Query(None),
    page: int = Query(1, ge=1),
    per_page: int = Query(DEFAULT_PAGE_SIZE, ge=1, le=MAX_PAGE_SIZE),
    db: Session = Depends(get_db),
current_admin: models.Admin = Depends(auth.get_current_admin),
):
    query = db.query(models.Renseignement)
    if direction:
        query = query.filter(models.Renseignement.direction.ilike(f"%{direction}%"))
    if statut:
        query = query.filter(models.Renseignement.statut == statut.lower())
    if cuid:
        query = query.filter(models.Renseignement.CUID.ilike(f"%{cuid}%"))
    if plateforme:
        query = query.filter(models.Renseignement.plateforme.ilike(f"%{plateforme}%"))

    total = query.count()
    records = query.order_by(models.Renseignement.id.desc()).offset((page - 1) * per_page).limit(per_page).all()

    return {
        "data": [
            {
                "id": r.id,
                "CUID": r.CUID,
                "Nom": r.Nom,
                "prenom": r.prenom,
                "statut": r.statut,
                "entreprise_partenaire": r.entreprise_partenaire or "",
                "direction": r.direction,
                "poste": r.poste or "",
                "referent": r.referent or "",
                "plateforme": r.plateforme or "",
                "description": r.description,
            }
            for r in records
        ],
        "total": total,
        "page": page,
        "per_page": per_page,
        "total_pages": max(1, (total + per_page - 1) // per_page),
    }


@router.get("/logs", summary="Logs des renseignements (soumissions)")
def logs_renseignements(
    cuid: Optional[str] = Query(None),
    adresse_mac: Optional[str] = Query(None),
    page: int = Query(1, ge=1),
    per_page: int = Query(DEFAULT_PAGE_SIZE, ge=1, le=MAX_PAGE_SIZE),
    db: Session = Depends(get_db),
current_admin: models.Admin = Depends(auth.get_current_admin),
):
    query = db.query(models.RenseignementLog)
    if cuid:
        query = query.filter(models.RenseignementLog.cuid.ilike(f"%{cuid}%"))
    if adresse_mac:
        query = query.filter(models.RenseignementLog.adresse_mac.ilike(f"%{adresse_mac}%"))

    total = query.count()
    records = query.order_by(models.RenseignementLog.id.desc()).offset((page - 1) * per_page).limit(per_page).all()

    return {
        "data": [
            {
                "id": r.id,
                "cuid": r.cuid,
                "adresse_ip": r.adresse_ip or "",
                "adresse_mac": r.adresse_mac or "",
                "date_creation": r.date_creation.isoformat() if r.date_creation else "",
            }
            for r in records
        ],
        "total": total,
        "page": page,
        "per_page": per_page,
        "total_pages": max(1, (total + per_page - 1) // per_page),
    }


@router.get("/logs/exporter", summary="Exporter les logs de renseignements en Excel")
def exporter_logs_renseignements(
    cuid: Optional[str] = Query(None),
    adresse_mac: Optional[str] = Query(None),
    db: Session = Depends(get_db),
current_admin: models.Admin = Depends(auth.get_current_admin),
):
    query = db.query(models.RenseignementLog)
    if cuid:
        query = query.filter(models.RenseignementLog.cuid.ilike(f"%{cuid}%"))
    if adresse_mac:
        query = query.filter(models.RenseignementLog.adresse_mac.ilike(f"%{adresse_mac}%"))

    records = query.order_by(models.RenseignementLog.id.desc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Logs Renseignements"

    BLEU_FONCE = "1F4E79"
    BLANC = "FFFFFF"
    en_tete_font = Font(bold=True, color=BLANC, size=11, name="Calibri")
    en_tete_fill = PatternFill(start_color=BLEU_FONCE, end_color=BLEU_FONCE, fill_type="solid")
    centre = Alignment(horizontal="center", vertical="center")
    bordure = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    colonnes = [
        ("ID", 6), ("CUID", 16), ("Adresse IP", 20), ("Adresse MAC", 22), ("Date de soumission", 25),
    ]
    for col_idx, (titre, largeur) in enumerate(colonnes, start=1):
        cell = ws.cell(row=1, column=col_idx, value=titre)
        cell.font = en_tete_font
        cell.fill = en_tete_fill
        cell.alignment = centre
        cell.border = bordure
        ws.column_dimensions[cell.column_letter].width = largeur

    ws.freeze_panes = "A2"

    for row_idx, rec in enumerate(records, start=2):
        valeurs = [
            rec.id, rec.cuid, rec.adresse_ip or "", rec.adresse_mac or "",
            rec.date_creation.strftime("%d/%m/%Y %H:%M:%S") if rec.date_creation else "",
        ]
        for col_idx, valeur in enumerate(valeurs, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=valeur)
            cell.border = bordure
            cell.alignment = centre

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"logs_renseignements_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/exporter", summary="Exporter en Excel")
def exporter_rapport(
    direction: Optional[str] = Query(None),
    statut: Optional[str] = Query(None),
    cuid: Optional[str] = Query(None),
    plateforme: Optional[str] = Query(None),
    db: Session = Depends(get_db),
current_admin: models.Admin = Depends(auth.get_current_admin),
):
    query = db.query(models.Renseignement)
    if direction:
        query = query.filter(models.Renseignement.direction.ilike(f"%{direction}%"))
    if statut:
        query = query.filter(models.Renseignement.statut == statut.lower())
    if cuid:
        query = query.filter(models.Renseignement.CUID.ilike(f"%{cuid}%"))
    if plateforme:
        query = query.filter(models.Renseignement.plateforme.ilike(f"%{plateforme}%"))

    records = query.order_by(models.Renseignement.id.desc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Renseignements"

    BLEU_FONCE = "1F4E79"
    BLEU_CLAIR = "D6E4F0"
    BLANC = "FFFFFF"
    en_tete_font = Font(bold=True, color=BLANC, size=11, name="Calibri")
    en_tete_fill = PatternFill(start_color=BLEU_FONCE, end_color=BLEU_FONCE, fill_type="solid")
    alt_fill = PatternFill(start_color=BLEU_CLAIR, end_color=BLEU_CLAIR, fill_type="solid")
    centre = Alignment(horizontal="center", vertical="center")
    gauche = Alignment(horizontal="left", vertical="center", wrap_text=True)
    bordure = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    colonnes = [
        ("ID", 6), ("CUID", 16), ("Nom", 22), ("Prénom", 22),
        ("Statut", 14), ("Entreprise partenaire", 22), ("Direction", 22),
        ("Poste", 22), ("Référent", 22),
        ("Plateforme", 25), ("Description", 45),
    ]

    for col_idx, (titre, largeur) in enumerate(colonnes, start=1):
        cell = ws.cell(row=1, column=col_idx, value=titre)
        cell.font = en_tete_font
        cell.fill = en_tete_fill
        cell.alignment = centre
        cell.border = bordure
        ws.column_dimensions[cell.column_letter].width = largeur

    ws.freeze_panes = "A2"

    for row_idx, rec in enumerate(records, start=2):
        valeurs = [
            rec.id, rec.CUID, rec.Nom, rec.prenom,
            rec.statut, rec.entreprise_partenaire or "", rec.direction,
            rec.poste or "", rec.referent or "",
            rec.plateforme or "", rec.description,
        ]
        fill = alt_fill if row_idx % 2 == 0 else None
        for col_idx, valeur in enumerate(valeurs, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=valeur)
            cell.border = bordure
            cell.alignment = centre if col_idx < 11 else gauche
            if fill:
                cell.fill = fill

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"renseignements_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.post(
    "/bloquer-compte",
    status_code=status.HTTP_201_CREATED,
    summary="Bloquer un compte après 3 tentatives sans description de plateforme renseignée",
)
def bloquer_compte(
    data: schemas.CompteBloqueCreate,
    db: Session = Depends(get_db),
    current_user: dict = Depends(auth.get_current_user),
):
    # Éviter les doublons si le blocage a déjà été enregistré pour ce CUID
    deja_bloque = db.query(models.BlockedAccount).filter(models.BlockedAccount.cuid == data.cuid).first()
    if deja_bloque:
        return {"message": "Ce compte est déjà bloqué."}

    plateformes_json = json.dumps(
        [{"text": p.text, "description": (p.description or "").strip()} for p in data.plateformes],
        ensure_ascii=False,
    )

    entry = models.BlockedAccount(
        cuid=data.cuid,
        nom=data.nom,
        prenom=data.prenom,
        direction=data.direction,
        statut=data.statut,
        entreprise_partenaire=data.entreprise_partenaire,
        plateformes=plateformes_json,
    )
    db.add(entry)
    db.commit()
    return {"message": "Compte bloqué : les informations ont été enregistrées."}


@router.get("/comptes-bloques", summary="Lister les comptes bloqués")
def lister_comptes_bloques(
    cuid: Optional[str] = Query(None),
    page: int = Query(1, ge=1),
    per_page: int = Query(DEFAULT_PAGE_SIZE, ge=1, le=MAX_PAGE_SIZE),
    db: Session = Depends(get_db),
):
    query = db.query(models.BlockedAccount)
    if cuid:
        query = query.filter(models.BlockedAccount.cuid.ilike(f"%{cuid}%"))

    total = query.count()
    records = query.order_by(models.BlockedAccount.id.desc()).offset((page - 1) * per_page).limit(per_page).all()

    def parse_plats(raw):
        try:
            return json.loads(raw) if raw else []
        except Exception:
            return []

    return {
        "data": [
            {
                "id": r.id,
                "cuid": r.cuid,
                "nom": r.nom or "",
                "prenom": r.prenom or "",
                "direction": r.direction or "",
                "statut": r.statut or "",
                "entreprise_partenaire": r.entreprise_partenaire or "",
                "plateformes": parse_plats(r.plateformes),
                "date_creation": r.created_at.isoformat() if r.created_at else "",
            }
            for r in records
        ],
        "total": total,
        "page": page,
        "per_page": per_page,
        "total_pages": max(1, (total + per_page - 1) // per_page),
    }


@router.delete("/comptes-bloques/{compte_id}", summary="Débloquer un compte utilisateur")
def debloquer_compte(
    compte_id: int,
    db: Session = Depends(get_db),
    current_admin: models.Admin = Depends(auth.get_current_admin),
):
    compte = db.query(models.BlockedAccount).filter(models.BlockedAccount.id == compte_id).first()
    if not compte:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Compte bloqué introuvable.")
    db.delete(compte)
    db.commit()
    return {"message": "Compte débloqué avec succès."}


@router.get("/comptes-bloques/exporter", summary="Exporter les comptes bloqués en Excel")
def exporter_comptes_bloques(
    cuid: Optional[str] = Query(None),
    db: Session = Depends(get_db),
):
    query = db.query(models.BlockedAccount)
    if cuid:
        query = query.filter(models.BlockedAccount.cuid.ilike(f"%{cuid}%"))

    records = query.order_by(models.BlockedAccount.id.desc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Comptes bloqués"

    ROUGE_FONCE = "7F1D1D"
    BLANC = "FFFFFF"
    en_tete_font = Font(bold=True, color=BLANC, size=11, name="Calibri")
    en_tete_fill = PatternFill(start_color=ROUGE_FONCE, end_color=ROUGE_FONCE, fill_type="solid")
    centre = Alignment(horizontal="center", vertical="center")
    gauche = Alignment(horizontal="left", vertical="center", wrap_text=True)
    bordure = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    colonnes = [
        ("ID", 6), ("CUID", 16), ("Nom", 20), ("Prénom", 20), ("Direction", 22),
        ("Statut", 14), ("Entreprise partenaire", 22), ("Plateformes sélectionnées", 55),
        ("Date de blocage", 22),
    ]
    for col_idx, (titre, largeur) in enumerate(colonnes, start=1):
        cell = ws.cell(row=1, column=col_idx, value=titre)
        cell.font = en_tete_font
        cell.fill = en_tete_fill
        cell.alignment = centre
        cell.border = bordure
        ws.column_dimensions[cell.column_letter].width = largeur

    ws.freeze_panes = "A2"

    for row_idx, rec in enumerate(records, start=2):
        try:
            plats = json.loads(rec.plateformes) if rec.plateformes else []
        except Exception:
            plats = []
        plats_txt = "; ".join(
            f"{p.get('text', '')}" + (f" ({p.get('description')})" if p.get("description") else " (description non renseignée)")
            for p in plats
        ) or "—"

        valeurs = [
            rec.id, rec.cuid, rec.nom or "", rec.prenom or "", rec.direction or "",
            rec.statut or "", rec.entreprise_partenaire or "", plats_txt,
            rec.created_at.strftime("%d/%m/%Y %H:%M:%S") if rec.created_at else "",
        ]
        for col_idx, valeur in enumerate(valeurs, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=valeur)
            cell.border = bordure
            cell.alignment = gauche if col_idx == 8 else centre

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"comptes_bloques_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
